import openpyxl
import json
import re

wb = openpyxl.load_workbook('ORION 1.0 .xlsx')
sheet = wb.active

teams = []
for r in range(2, sheet.max_row + 1):
    raw_team = sheet.cell(row=r, column=3).value
    raw_leader = sheet.cell(row=r, column=5).value
    raw_email = sheet.cell(row=r, column=6).value
    raw_dept = sheet.cell(row=r, column=7).value
    raw_year = sheet.cell(row=r, column=8).value
    raw_college = sheet.cell(row=r, column=9).value
    raw_ps = sheet.cell(row=r, column=35).value
    
    if raw_team is None or str(raw_team).strip() == '':
        continue

    u_clean = re.sub(r'\s+', '', str(raw_team)).upper()
    p_clean = re.sub(r'[^A-Z0-9]', '', str(raw_leader or '').upper())
    
    members = []
    if raw_leader:
        members.append({
            "name": str(raw_leader).strip(),
            "email": str(raw_email or '').strip(),
            "department": str(raw_dept or '').strip(),
            "year": str(raw_year or '').strip(),
            "institution": str(raw_college or '').strip(),
            "isLeader": True
        })
    
    for m_idx, col_start in enumerate([10, 15, 20, 25, 30], start=2):
        m_name = sheet.cell(row=r, column=col_start).value
        m_email = sheet.cell(row=r, column=col_start+1).value
        m_dept = sheet.cell(row=r, column=col_start+2).value
        m_year = sheet.cell(row=r, column=col_start+3).value
        m_inst = sheet.cell(row=r, column=col_start+4).value
        if m_name and str(m_name).strip():
            members.append({
                "name": str(m_name).strip(),
                "email": str(m_email or '').strip(),
                "department": str(m_dept or '').strip(),
                "year": str(m_year or '').strip(),
                "institution": str(m_inst or '').strip(),
                "isLeader": False
            })

    reg_id = str(sheet.cell(row=r, column=2).value or f'ORION-S{r:04d}').strip()

    teams.append({
        "reg_id": reg_id,
        "team_name": str(raw_team).strip(),
        "username": u_clean,
        "password": p_clean,
        "leader_name": str(raw_leader or '').strip(),
        "leader_email": str(raw_email or '').strip(),
        "institution": str(raw_college or '').strip(),
        "department": str(raw_dept or '').strip(),
        "year": str(raw_year or '').strip(),
        "problem_statement": str(raw_ps or 'ORION-PS-04').strip(),
        "members": members
    })

with open('.data/extracted_teams.json', 'w', encoding='utf-8') as f:
    json.dump(teams, f, indent=2)

print(f"Extracted {len(teams)} teams to .data/extracted_teams.json")
